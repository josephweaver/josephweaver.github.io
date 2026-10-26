"""Build an editable half-letter rulebook and a duplex US Letter booklet.

Requires desktop Microsoft Word on Windows for PDF export:
  python -m pip install python-docx markdown-it-py pypdf pywin32
  python fiction/24system/tools/build_playtest_booklet.py

The Markdown remains the source of truth. Rebuild after changing the rules.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt
from pypdf import PdfReader, PdfWriter, Transformation


RULES_DIR = Path(__file__).resolve().parents[1]
PAGE_WIDTH = 5.5
PAGE_HEIGHT = 8.5
MARGIN = 0.45
TEXT_WIDTH = PAGE_WIDTH - 2 * MARGIN
INK = RGBColor.from_string("202020")


def xml_element(name, **attributes):
    element = OxmlElement(name)
    for key, value in attributes.items():
        element.set(qn(key), str(value))
    return element


def set_columns(section, count):
    columns = section._sectPr.find(qn("w:cols"))
    if columns is None:
        columns = xml_element("w:cols")
        section._sectPr.append(columns)
    columns.set(qn("w:num"), str(count))
    columns.set(qn("w:space"), "230")
    columns.set(qn("w:equalWidth"), "1")


def change_columns(document, count, *, force=False):
    current = document.sections[-1]._sectPr.find(qn("w:cols"))
    if not force and current is not None and current.get(qn("w:num"), "1") == str(count):
        return
    section = document.add_section(WD_SECTION_START.CONTINUOUS)
    set_columns(section, count)
    # Keep layout-only paragraphs from consuming a line of body text.
    boundary = document.paragraphs[-1]
    boundary.paragraph_format.space_before = Pt(0)
    boundary.paragraph_format.space_after = Pt(0)
    boundary.paragraph_format.line_spacing = Pt(1)
    boundary.add_run().font.size = Pt(1)


def configure(document, title):
    section = document.sections[0]
    section.page_width = Inches(PAGE_WIDTH)
    section.page_height = Inches(PAGE_HEIGHT)
    section.left_margin = section.right_margin = Inches(MARGIN)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.22)
    set_columns(section, 1)

    normal = document.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    for style_name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 11), ("Heading 3", 10)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    document.styles["Heading 1"].paragraph_format.page_break_before = True
    document.styles["Title"].paragraph_format.space_before = Pt(0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("24-SYSTEM  /  AUGUST 2026 PLAYTEST")
    run.font.name = "Aptos"
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string("666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    field = xml_element("w:fldSimple", **{"w:instr": "PAGE"})
    field.append(run._r)
    footer._p.append(field)

    document.core_properties.title = title
    document.core_properties.subject = "Half-letter playtest rulebook with two-column text and full-width tables"
    document.core_properties.author = "Joseph Weaver"
    document.core_properties.keywords = "24-System, playtest, booklet"


def link_target(href):
    if ":" in href or href.startswith("#"):
        return href
    # Keep source references useful in the exported PDF and in copied Word files.
    return "https://github.com/josephweaver/josephweaver.github.io/blob/main/fiction/24system/" + quote(href, safe="/#")


def write_inline(paragraph, children, *, table=False, header=False, font_size=8.5):
    bold = italic = False
    hyperlink = None
    for token in children or []:
        if token.type == "strong_open":
            bold = True
            continue
        if token.type == "strong_close":
            bold = False
            continue
        if token.type == "em_open":
            italic = True
            continue
        if token.type == "em_close":
            italic = False
            continue
        if token.type == "link_open":
            from docx.opc.constants import RELATIONSHIP_TYPE

            hyperlink = xml_element("w:hyperlink")
            relationship = paragraph.part.relate_to(link_target(token.attrGet("href")), RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            hyperlink.set(qn("r:id"), relationship)
            paragraph._p.append(hyperlink)
            continue
        if token.type == "link_close":
            hyperlink = None
            continue
        if token.type not in ("text", "code_inline", "softbreak", "hardbreak"):
            raise ValueError(f"Unsupported inline Markdown: {token.type}")
        run = paragraph.add_run(" " if token.type == "softbreak" else "\n" if token.type == "hardbreak" else token.content)
        run.bold = bold or header
        run.italic = italic
        if table:
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
        if token.type == "code_inline":
            run.font.name = "Consolas"
            run.font.size = Pt(font_size if table else 9)
        if hyperlink is not None:
            run.font.color.rgb = INK
            run.underline = True
            hyperlink.append(run._r)


def table_spec(headers):
    """Widths in inches, tuned for the source's wide reference tables."""
    first = headers[0]
    if first == "Armor":
        return [1.03, 0.45, 0.95, 0.53, 0.63, 0.62, 0.39], 7.5
    if first == "Fear":
        return [0.32, 0.73, 0.82, 0.85, 1.88], 8
    if first == "Pack":
        return [0.55, 0.55, 0.45, 3.05], 8.5
    if first == "Location":
        return [0.73, 3.87], 8.5
    if first == "Final pushed result":
        return [1.13, 3.47], 8.5
    if first == "Individual skill die":
        return [1.22, 2.28, 1.10], 8.5
    if first == "Rank":
        return [0.38, 2.15, 1.07, 1.0], 8.5
    if first == "Skill":
        return [0.73, 3.87], 8.5
    if first == "Difficulty":
        return [0.78, 0.30, 3.52], 8.5
    if first == "Roll":
        return [0.85, 3.75], 8.5
    if first == "Item":
        return [2.63, 1.14, 0.83], 8.5
    return [TEXT_WIDTH / len(headers)] * len(headers), 8.5


def add_table(document, tokens):
    rows = []
    row = None
    cell = None
    for token in tokens:
        if token.type == "tr_open":
            row = []
        elif token.type in ("th_open", "td_open"):
            cell = {"header": token.type == "th_open", "alignment": token.attrGet("style") or "", "children": []}
        elif token.type == "inline":
            cell["children"] = token.children
        elif token.type in ("th_close", "td_close"):
            row.append(cell)
        elif token.type == "tr_close":
            rows.append(row)
    headers = ["".join(t.content for t in c["children"] if t.type in ("text", "code_inline")) for c in rows[0]]
    widths, font_size = table_spec(headers)
    change_columns(document, 1)
    table = document.add_table(rows=0, cols=len(headers))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(round(TEXT_WIDTH * 1440)))
    margins = xml_element("w:tblCellMar")
    for side, value in (("top", 35), ("bottom", 35), ("left", 45), ("right", 45)):
        margins.append(xml_element(f"w:{side}", **{"w:w": value, "w:type": "dxa"}))
    properties.append(margins)
    for row_index, data in enumerate(rows):
        row = table.add_row()
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(xml_element("w:cantSplit"))
        if row_index == 0:
            row_properties.append(xml_element("w:tblHeader"))
        for cell, content, width in zip(row.cells, data, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.02
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = row_index == 0
            if "right" in content["alignment"]:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif "center" in content["alignment"]:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            write_inline(paragraph, content["children"], table=True, header=row_index == 0, font_size=font_size)
            if row_index == 0 or row_index % 2 == 0:
                cell._tc.get_or_add_tcPr().append(xml_element("w:shd", **{"w:fill": "DDDDDD" if row_index == 0 else "F4F4F4"}))
    # Word needs a paragraph between a table and its section boundary.
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.add_run().font.size = Pt(1)
    change_columns(document, 2)


def build_docx(source, output):
    tokens = MarkdownIt("commonmark").enable("table").parse(source.read_text(encoding="utf-8"))
    title = next(t.content for t in tokens if t.type == "inline")
    document = Document()
    configure(document, title)
    lists = []
    index = 0
    first_chapter = True
    while index < len(tokens):
        token = tokens[index]
        if token.type == "table_open":
            end = index + 1
            while tokens[end].type != "table_close":
                end += 1
            add_table(document, tokens[index:end + 1])
            index = end + 1
            continue
        if token.type in ("bullet_list_open", "ordered_list_open"):
            lists.append({"ordered": token.type == "ordered_list_open", "next": int(token.attrGet("start") or 1), "fresh": False})
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif token.type == "list_item_open":
            lists[-1]["fresh"] = True
        elif token.type in ("heading_open", "paragraph_open"):
            inline = tokens[index + 1]
            if inline.type != "inline":
                raise ValueError(f"Expected inline Markdown after {token.type}")
            if token.type == "heading_open":
                level = int(token.tag[1:])
                if level > 1:
                    change_columns(document, 2, force=level == 2 and not first_chapter)
                paragraph = document.add_paragraph(style="Title" if level == 1 else f"Heading {min(level - 1, 3)}")
                if level == 2 and first_chapter:
                    paragraph.paragraph_format.page_break_before = False
                    first_chapter = False
            else:
                paragraph = document.add_paragraph()
                if lists:
                    depth = len(lists) - 1
                    paragraph.paragraph_format.left_indent = Inches(0.11 + depth * 0.13)
                    paragraph.paragraph_format.first_line_indent = Inches(-0.11)
                    paragraph.paragraph_format.space_after = Pt(2.5)
                    if lists[-1]["fresh"]:
                        marker = f"{lists[-1]['next']}. " if lists[-1]["ordered"] else "• "
                        paragraph.add_run(marker)
                        lists[-1]["next"] += 1
                        lists[-1]["fresh"] = False
                if len(inline.children or []) >= 3 and inline.children[0].type == "strong_open" and inline.children[-1].type == "strong_close":
                    paragraph.paragraph_format.keep_with_next = True
                    paragraph.paragraph_format.space_before = Pt(4)
            write_inline(paragraph, inline.children)
            index += 2
        elif token.type not in ("heading_close", "paragraph_close", "list_item_close"):
            raise ValueError(f"Unsupported block Markdown: {token.type}")
        index += 1
    # Balance the last page's text across both columns.
    change_columns(document, 1)
    document.save(output)
    return tokens


def export_pdf(docx_path, pdf_path):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        word.AutomationSecurity = 3
        document = word.Documents.Open(str(docx_path), ReadOnly=True, AddToRecentFiles=False, Visible=False)
        document.Fields.Update()
        document.Repaginate()
        document.ExportAsFixedFormat(str(pdf_path), 17, OpenAfterExport=False, OptimizeFor=0, CreateBookmarks=1, DocStructureTags=True)
    finally:
        if document is not None:
            document.Close(False)
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()


def impose_booklet(reader_path, booklet_path):
    reader = PdfReader(reader_path)
    count = len(reader.pages)
    padded = (count + 3) // 4 * 4
    writer = PdfWriter()
    pairings = []
    for sheet in range(padded // 4):
        pairs = ((padded - 1 - sheet * 2, sheet * 2), (sheet * 2 + 1, padded - 2 - sheet * 2))
        for left, right in pairs:
            spread = writer.add_blank_page(width=792, height=612)
            for offset, page_index in ((0, left), (396, right)):
                if page_index < count:
                    page = reader.pages[page_index]
                    if abs(float(page.mediabox.width) - 396) > 0.1 or abs(float(page.mediabox.height) - 612) > 0.1:
                        raise ValueError("Expected half-letter reader pages (396 × 612 PDF points)")
                    spread.merge_transformed_page(page, Transformation().translate(tx=offset), expand=False)
            pairings.append([left + 1 if left < count else None, right + 1 if right < count else None])
    writer.add_metadata({"/Title": "24-System: August 2026 Playtest — Print Booklet", "/Subject": "US Letter landscape; actual size; duplex short edge; one PDF page per sheet side"})
    preferences = writer.create_viewer_preferences()
    preferences.duplex = "/DuplexFlipShortEdge"
    preferences.print_scaling = "/None"
    writer.write(booklet_path)
    return {"reader_pages": count, "booklet_pages": padded, "blank_padding_pages": padded - count, "sheets": padded // 4, "sheet_side_page_pairs": pairings}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=RULES_DIR / "2026-08-playtest.md")
    parser.add_argument("--output-dir", type=Path, default=RULES_DIR / "print")
    arguments = parser.parse_args()
    source = arguments.source.resolve()
    destination = arguments.output_dir.resolve()
    destination.mkdir(parents=True, exist_ok=True)
    base = source.stem
    docx_path = destination / f"{base}.docx"
    reader_path = destination / f"{base}-reader.pdf"
    booklet_path = destination / f"{base}-booklet.pdf"
    tokens = build_docx(source, docx_path)
    print(f"Created {docx_path}", flush=True)
    export_pdf(docx_path, reader_path)
    print(f"Created {reader_path}", flush=True)
    details = impose_booklet(reader_path, booklet_path)
    details["tables"] = sum(t.type == "table_open" for t in tokens)
    (destination / f"{base}-print-info.json").write_text(json.dumps(details, indent=2) + "\n", encoding="utf-8")
    print(f"Created {booklet_path}", flush=True)
    print(json.dumps(details, indent=2))


if __name__ == "__main__":
    main()
